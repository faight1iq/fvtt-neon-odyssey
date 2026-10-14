"""Neon Odyssey starship components v2: build compendium items from the Design Studio workbook.

Source: tools/design-studio-xlsx.json (the Google Sheet "Neon Odyssey Starship Design Studio",
exported as .xlsx and wrapped in base64). Output: tools/starship-components-v2.json, which the
Foundry import (tools/import-starship-components-v2.js, run in the browser console) loads into
the "Starship Components" compendium.

Every component carries flags["fvtt-neon-odyssey"].component = {
  v: 2, tab, id, slot, size, stats{...stat block...}, info{...item columns...}, mishaps[],
  hudSystem, stationRoles[]
}
The ship calculator (scripts/starship-designer.mjs) sums `stats` exactly like the sheet's
Design Analysis tab.
"""
import base64, hashlib, html, io, json, os, re, sys
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "design-studio-xlsx.json")
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, "starship-components-v2.json")
MOD = "fvtt-neon-odyssey"

if SRC.endswith(".xlsx"):
    raw = {}; wb = openpyxl.load_workbook(SRC, data_only=True)
else:
    raw = json.load(open(SRC))
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(raw["xlsx"])), data_only=True)

def hid(*parts):
    """Deterministic 16-character Foundry id."""
    h = int(hashlib.sha1("|".join(map(str, parts)).encode()).hexdigest(), 16)
    a = "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789"; s = ""
    while len(s) < 16: s += a[h % 62]; h //= 62
    return s

def num(v):
    if v is None or v == "": return 0
    if isinstance(v, (int, float)): return int(v) if float(v).is_integer() else float(v)
    t = str(v).replace(",", "").replace("+", "").strip()
    try:
        f = float(t); return int(f) if f.is_integer() else f
    except ValueError: return 0

def text(v):
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

# Stat block names (as in the sheet header) -> keys used in Foundry.
STAT_KEYS = {
    "AC": "ac", "Hit Die": "hitDie", "Hit Dice": "hitDice", "HP Bonus": "hpBonus", "MT": "mt",
    "Speed": "speed", "STR": "str", "DEX": "dex", "CON": "con", "INT": "int", "WIS": "wis", "CHA": "cha",
    "Acrobatics": "acrobatics", "Athletics": "athletics", "Deception": "deception",
    "Perception": "perception", "Stealth": "stealth", "Cargo": "cargo", "Crew": "crew",
    "Passengers": "passengers", "Power": "power", "Compartments": "compartments", "Cost": "cost",
    "Berths": "berths", "Shield AC": "shieldAc", "Shield Charges": "shieldCharges",
    "Ablative HP": "ablativeHp", "Ablative Regen": "ablativeRegen", "Prisoners": "prisoners"
}

def stat_block(ws, row, header_row, first_col, last_col):
    """Read a stat block like the sheet's MATCH(header): the first column with each name wins."""
    out = {}
    for c in range(first_col, last_col + 1):
        name = ws.cell(header_row, c).value
        key = STAT_KEYS.get(text(name))
        if not key or key in out: continue
        out[key] = num(ws.cell(row, c).value)
    return {k: v for k, v in out.items() if v}

def col_index(letter):
    return openpyxl.utils.column_index_from_string(letter)

def item_columns(ws, row, header_row, last_col):
    """The descriptive columns (A..S) of a component row, by header name."""
    info = {}
    for c in range(1, last_col + 1):
        h = text(ws.cell(header_row, c).value)
        v = ws.cell(row, c).value
        if h and v not in (None, ""): info[h] = v if isinstance(v, (int, float)) else text(v)
    return info

def mishap_list(s):
    return [m.strip() for m in text(s).split(",") if m.strip()]

# ---------------------------------------------------------------- HUD systems and crew roles
HUD_SYSTEM = {
    "shipClass": "hull", "material": "hull", "descriptor": None, "shipType": None,
    "propulsion": "propulsion", "power": "power", "armor": "armor", "deflector": "shields",
    "ablative": "shields", "sensory": "navigation", "nexus": "command", "comms": "command",
    "habitat": "habitat", "weapon": "weapons"
}

# Station components (Habitats tab) -> crew cards (John approved 2026-10-12).
STATION_ROLES = [
    (r"^Saddle$", ["pilot"]),
    (r"^Cockpit$", ["pilot"]),
    (r"^Tandem Cockpit$", ["pilot", "copilot"]),
    (r"^Cabin Seating", ["pilot"]),
    (r"^Command & Control", ["commander"]),
    (r"^Auxiliary Bridge", []),
    (r"^Helm Station", ["helmsman"]),
    (r"^Gunnery Station", ["gunner"]),
    (r"^Engineering Station", ["engineer"]),
    (r"^Sensors Station", ["navigator"]),
    (r"^Electronic Warfare Station", ["navigator"]),
    (r"^Damage Control Station", ["engineer"]),
    (r"^Communications Station", ["commander"]),
    (r"^Tactical Station", ["navigator"]),
    (r"^Flight Control Station", ["helmsman"]),
]

def station_roles(name):
    for pat, roles in STATION_ROLES:
        if re.search(pat, name): return roles
    return None

# Habitat stations that sit in the command and navigation boxes of the HUD schematic.
def habitat_system(name):
    if re.search(r"^(Command & Control|Auxiliary Bridge|Communications Station|Cockpit|Tandem Cockpit|Saddle|Cabin Seating)", name): return "command"
    if re.search(r"^(Helm Station|Flight Control Station)", name): return "command"
    if re.search(r"^(Sensors Station|Electronic Warfare Station|Tactical Station)", name): return "navigation"
    if re.search(r"^Gunnery Station", name): return "weapons"
    if re.search(r"^(Engineering Station|Damage Control Station)", name): return "power"
    return "habitat"

HABITAT_FOLDERS = {1: "Seating & Cockpits", 2: "Bridge & Stations", 3: "Crew Quarters", 4: "Passenger Space",
                   5: "Amenities", 6: "Services", 7: "Large Habitats", 8: "Bays, Docks & Hangars", 9: "Ship Systems & Holds"}

ICONS = {
    "shipClass": "icons/svg/cube.svg", "material": "icons/svg/statue.svg", "descriptor": "icons/svg/aura.svg",
    "shipType": "icons/svg/anchor.svg", "propulsion": "icons/svg/wing.svg", "power": "icons/svg/lightning.svg",
    "armor": "icons/svg/shield.svg", "deflector": "icons/svg/mage-shield.svg", "ablative": "icons/svg/ice-shield.svg",
    "sensory": "icons/svg/eye.svg", "nexus": "icons/svg/clockwork.svg", "comms": "icons/svg/sound.svg",
    "habitat": "icons/svg/house.svg", "weapon": "icons/svg/target.svg"
}
STATION_ICON = "icons/svg/village.svg"

FOLDER_NAMES = {
    "shipClass": "Ship Classes", "material": "Materials", "descriptor": "Descriptors", "shipType": "Ship Types",
    "propulsion": "Propulsion", "power": "Power Sources", "armor": "Armor", "deflector": "Shields",
    "ablative": "Shields", "sensory": "Sensory Gear", "nexus": "Nexus Cores", "comms": "Communications",
    "habitat": "Habitats", "weapon": "Weapons"
}

# ---------------------------------------------------------------- item builders
SIZES = ["Tiny", "Small", "Medium", "Large", "Huge", "Gargantuan"]
DND_DAMAGE = {"acid", "bludgeoning", "cold", "fire", "force", "lightning", "necrotic", "piercing",
              "poison", "psychic", "radiant", "slashing", "thunder"}

def describe(name, info, stats, mishaps, extra=""):
    rows = []
    for k, v in info.items():
        if k in ("ID", "Desc Name", "Mishap List", "Description", "Name") or k.startswith("Compartment"): continue
        rows.append(f"<tr><th>{html.escape(k)}</th><td>{html.escape(text(v))}</td></tr>")
    srow = [f"<tr><th>{html.escape(k)}</th><td>{v:+g}</td></tr>" if isinstance(v, (int, float)) and k not in ("cost", "hitDie", "hitDice") else f"<tr><th>{html.escape(k)}</th><td>{v}</td></tr>" for k, v in stats.items()]
    desc = text(info.get("Description", ""))
    parts = [f"<p>{html.escape(desc)}</p>" if desc else ""]
    if extra: parts.append(extra)
    if rows: parts.append("<table class='no-comp'><tbody>" + "".join(rows) + "</tbody></table>")
    if srow: parts.append("<p><strong>Ship stat effects</strong></p><table class='no-comp'><tbody>" + "".join(srow) + "</tbody></table>")
    if mishaps: parts.append(f"<p><strong>Possible Mishaps:</strong> {html.escape(', '.join(mishaps))}</p>")
    parts.append("<p><em>Neon Odyssey Starship Design Studio component (draft).</em></p>")
    return "".join(parts)

def base_item(tab, slot, cid, name, stats, info, mishaps, size=None, extra_desc="", folder_sub=None, roles=None, system=None):
    iid = hid("v2", tab, cid)
    comp = {"v": 2, "tab": tab, "id": cid, "slot": slot, "stats": stats, "info": info, "mishaps": mishaps,
            "hudSystem": system if system is not None else HUD_SYSTEM.get(slot)}
    if size: comp["size"] = size
    if roles is not None: comp["stationRoles"] = roles
    cost = num(stats.get("cost", 0))
    return {
        "_id": iid, "name": name, "type": "equipment",
        "img": STATION_ICON if roles else ICONS[slot],
        "system": {
            "description": {"value": describe(name, info, stats, mishaps, extra_desc), "chat": ""},
            "source": {"custom": "Neon Odyssey Starship Design Studio", "revision": 1, "rules": "2024"},
            "identifier": re.sub(r"[^a-z0-9]+", "-", f"{cid}".lower()).strip("-"),
            "quantity": 1, "price": {"value": cost, "denomination": "gp"}, "weight": {"value": 0, "units": "lb"},
            "type": {"value": "vehicle", "baseItem": ""}, "equipped": True
        },
        "flags": {MOD: {"component": comp}},
        "_folder": [FOLDER_NAMES[slot]] + ([folder_sub] if folder_sub else []),
        "sort": 0
    }

def weapon_item(cid, name, stats, info, mishaps):
    # Compendium folders (John, 2026-10-12): missiles in Weapons/Missiles, torpedo warheads in
    # Weapons/Warheads; everything else by damage technology.
    delivery = text(info.get("Delivery"))
    sub = {"Missile": "Missiles", "Warhead": "Warheads", "Launcher": "Torpedo Launchers"}.get(delivery) or (text(info.get("Energy")) or None)
    it = base_item("Weapons", "weapon", cid, name, stats, info, mishaps, size=text(info.get("Size")), folder_sub=sub)
    it["type"] = "weapon"
    dmg = text(info.get("Damage")); dtype = text(info.get("Damage Type")).lower()
    m = re.match(r"^(\d*)d(\d+)$", dmg)
    sysd = it["system"]
    sysd.pop("type", None); sysd.pop("equipped", None)
    sysd.update({
        "type": {"value": "siege", "baseItem": ""},
        "range": {"value": None, "long": None, "units": "ft", "reach": None},
        "damage": {"base": {"number": int(m.group(1) or 1) if m else None, "denomination": int(m.group(2)) if m else None,
                            "types": [dtype] if dtype in DND_DAMAGE else [], "custom": {"enabled": False},
                            "bonus": "", "scaling": {"number": 1}},
                   "versatile": {"types": [], "custom": {"enabled": False}, "scaling": {"number": 1}}},
        "properties": [], "equipped": True, "proficient": 1
    })
    act_id = hid("v2act", cid)
    sysd["activities"] = {act_id: {
        "_id": act_id, "type": "attack", "sort": 0,
        "activation": {"type": "action", "value": 1, "override": False, "condition": ""},
        "attack": {"ability": "none", "bonus": "", "flat": False, "critical": {"threshold": None},
                   "type": {"value": "ranged", "classification": "weapon"}},
        "damage": {"critical": {"bonus": ""}, "includeBase": True, "parts": []},
        "range": {"value": "", "units": "self", "special": text(info.get("Range")), "override": False},
        "consumption": {"scaling": {"allowed": False}, "spellSlot": True, "targets": []}, "description": {},
        "duration": {"units": "inst", "concentration": False, "override": False}, "effects": [], "flags": {},
        "target": {"template": {"contiguous": False, "stationary": False, "units": "ft"}, "affects": {"choice": False}, "override": False, "prompt": True},
        "uses": {"spent": 0, "recovery": []}
    }} if m else {}
    flags = it["flags"][MOD]
    flags["rangeBand"] = text(info.get("Range"))
    flags["crew"] = num(info.get("Crew"))
    return it

items = []

def read_tab(tab, slot, header_row=1, stat_first="T", stat_last="AV", info_last="S", id_col=1, name_col=2,
             size_col=None, extra=None):
    ws = wb[tab]
    for r in range(header_row + 1, ws.max_row + 1):
        cid = text(ws.cell(r, id_col).value)
        name = text(ws.cell(r, name_col).value)
        if not cid or not name: continue
        stats = stat_block(ws, r, header_row, col_index(stat_first), min(col_index(stat_last), ws.max_column))
        info = item_columns(ws, r, header_row, col_index(info_last))
        mishaps = mishap_list(info.get("Mishap List", ""))
        if slot == "weapon":
            items.append(weapon_item(cid, name, stats, info, mishaps)); continue
        size = text(ws.cell(r, size_col).value) if size_col else None
        sub = None; roles = None; system = None
        if slot == "habitat":
            n = int(re.match(r"(\d+)", cid).group(1)) // 100
            sub = HABITAT_FOLDERS.get(n)
            roles = station_roles(name)
            system = habitat_system(name)
        if slot in ("deflector", "ablative"):
            slot = "deflector" if text(info.get("Type")) == "Deflector" else "ablative"
        it = base_item(tab, slot, cid, name, stats, info, mishaps, size=size, folder_sub=sub, roles=roles, system=system)
        if extra: extra(it, info)
        items.append(it)
        slot = {"deflector": "deflector", "ablative": "deflector"}.get(slot, slot) if tab == "Shields" else slot

# Ship Classes: header row 6, stat block L..AH; ID in A, name in B.
read_tab("Ship Classes", "shipClass", header_row=6, stat_first="L", stat_last="AH", info_last="K")
# Materials: header row 2, stat block K..AG.
read_tab("Materials", "material", header_row=2, stat_first="K", stat_last="AG", info_last="J")
read_tab("Propulsion", "propulsion")
read_tab("Power Source", "power")
read_tab("Armor", "armor")
read_tab("Shields", "deflector", size_col=5)
read_tab("Sensory Gear", "sensory")
read_tab("Nexus Core", "nexus")
read_tab("Communications", "comms")
read_tab("Habitats", "habitat")
read_tab("Weapons", "weapon")

# Descriptors: the Design Analysis reads Speed from D, AC from E, HP Bonus from F, MT from G (DT),
# Power from I. Key = the label in M ("50 Speedy").
ws = wb["Descriptors"]
for r in range(2, ws.max_row + 1):
    key = text(ws.cell(r, 13).value); name = text(ws.cell(r, 2).value)
    if not key or not name: continue
    stats = {k: num(ws.cell(r, c).value) for k, c in (("speed", 4), ("ac", 5), ("hpPct", 6), ("mt", 7), ("power", 9))}  # HP is a % of hull HP (John, 2026-10-12)
    stats = {k: v for k, v in stats.items() if v}
    info = {"Descriptor": name, "Mishap": text(ws.cell(r, 8).value), "Cost": text(ws.cell(r, 10).value)}
    info = {k: v for k, v in info.items() if v}
    items.append(base_item("Descriptors", "descriptor", key, name, stats, info, []))

# Ship Type: mishap list only. Key = label in M ("40 Starship").
ws = wb["Ship Type"]
for r in range(2, ws.max_row + 1):
    key = text(ws.cell(r, 13).value); name = text(ws.cell(r, 2).value)
    if not key or not name: continue
    items.append(base_item("Ship Type", "shipType", key, name, {}, {"Type": name}, mishap_list(ws.cell(r, 4).value)))

# Fix the Shields slot per row (Deflector vs Ablative) after the fact.
for it in items:
    c = it["flags"][MOD]["component"]
    if c["tab"] == "Shields":
        c["slot"] = "deflector" if text(c["info"].get("Type")) == "Deflector" else "ablative"
        it["img"] = ICONS[c["slot"]]

# ---------------------------------------------------------------- rule overrides (John, 2026-10-12)
# Applied on import until the sheet carries the same values.
#  - Propulsion and Power Source take 12.5% of the ship's compartments each (25% together, was 30%).
#  - Fractional compartments elsewhere (Neutrino sensors/comms, Warded Cladding, ...) are 1 compartment.
#  - Brig and Detention Block hold prisoners (Brig 4, Detention Block 8), not passengers. Their crew
#    is 0.5 per block for all brig blocks together, rounded down (the calculator applies it).
OVERRIDES = []
for it in items:
    c = it["flags"][MOD]["component"]; st = c["stats"]
    v = st.get("compartments", 0)
    if c["slot"] in ("propulsion", "power") and isinstance(v, float) and -1 < v < 0:
        st["compartments"] = -0.125; OVERRIDES.append((c["id"], "compartments", v, -0.125))
    elif isinstance(v, float) and -1 < v < 0:
        st["compartments"] = -1; OVERRIDES.append((c["id"], "compartments", v, -1))
    if c["slot"] == "habitat" and re.match(r"^62[12]-", c["id"]):
        cap = 8 if "Detention" in it["name"] else 4
        st.pop("passengers", None); st.pop("crew", None); st["prisoners"] = cap
        c["brig"] = True
        OVERRIDES.append((c["id"], "prisoners", None, cap))
    # Ship class compartment budget: 2.1 x Tonns^(1/3) (sheet: 2 x), so a full Large fits.
    if c["slot"] == "shipClass" and num(c["info"].get("Tonns")):
        nb = round(2.1 * num(c["info"]["Tonns"]) ** (1 / 3))
        OVERRIDES.append((c["id"], "compartments", st.get("compartments"), nb)); st["compartments"] = nb
    # Helms (spelljamming power sources) are worked by the pilot: their station is the Helmsman's.
    if c["slot"] == "power" and re.search(r"Helm|Orbus", it["name"]):
        c["systemRole"] = "helmsman"
    # Galley and Medical Bay are worked by a Cook and a Medic (crew cards), unless automated.
    if c["slot"] == "habitat" and c["id"] in ("501-Galley / Mess Hall", "601-Medical Bay"):
        c["seatRole"] = "cook" if c["id"].startswith("501") else "medic"
    # Descriptions follow the stats.
    it["system"]["description"]["value"] = describe(it["name"], c["info"], st, c["mishaps"])

# Ammo (John, 2026-10-12): Ammunition Magazines come in the six sizes; each holds 12 units of its own
# size (a missile uses 6, a ballistic round 1). Missiles and Ballistic weapons use ammo; MLA ammo is
# too small to track.
for it in items:
    c = it["flags"][MOD]["component"]
    m = re.match(r"^Ammunition Magazine \((Tiny|Small|Medium|Large|Huge|Gargantuan)\)$", it["name"])
    if c["slot"] == "habitat" and m:
        c["magazine"] = {"size": m.group(1)}
    if c["slot"] == "weapon":
        info = c.get("info", {})
        missile = info.get("Delivery") == "Missile"
        ballistic = info.get("Energy") == "Ballistic"
        if (missile or ballistic) and info.get("Size"):
            c["ammo"] = {"size": info["Size"], "cost": 6 if missile else 1}

# Torpedo launchers (John, 2026-10-12): Tubes hold 1 loaded torpedo, Racks 4 (one salvo). Sized by
# torpedo class; the smallest ship that can mount each is in its description ("Mount on a X ship").
for it in items:
    c = it["flags"][MOD]["component"]
    if c["slot"] == "weapon" and c.get("info", {}).get("Delivery") == "Launcher":
        m = re.search(r"Mount on an? (\w+) ship", c["info"].get("Description", ""))
        c["launcher"] = {"class": c["info"].get("Size"), "tubes": 4 if "Rack" in it["name"] else 1,
                         "minShip": m.group(1) if m else "Tiny"}

# Power draw scales with size (John, 2026-10-12): weapons and shields draw power per compartment.
# The rate per technology is its Small model's draw per compartment (the smallest size listed if
# there is no Small one). The calculator multiplies it by the part's compartments.
SIZE_ORDER = ["Tiny", "Small", "Medium", "Large", "Huge", "Gargantuan"]
fam = {}
for it in items:
    c = it["flags"][MOD]["component"]
    if c["slot"] not in ("weapon", "deflector", "ablative"): continue
    m = re.match(r"^(.*) \(([^()]+)\)$", it["name"])
    if m and m.group(2) in SIZE_ORDER: fam.setdefault(m.group(1), {})[m.group(2)] = c
for family, sizes in fam.items():
    src = sizes.get("Small") or next((sizes[z] for z in SIZE_ORDER if z in sizes), None)
    comp = -num(src["stats"].get("compartments")) if src else 0
    rate = (-num(src["stats"].get("power")) / comp) if comp > 0 else 0
    for c in sizes.values(): c["powerRate"] = round(rate, 4)

# Mishaps tab (reference text) for the import log.
mishap_ref = []
ws = wb["Mishaps"]
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 1).value: mishap_ref.append({"id": text(ws.cell(r, 1).value), "name": text(ws.cell(r, 2).value),
        "group": text(ws.cell(r, 3).value), "usedBy": text(ws.cell(r, 5).value), "reference": text(ws.cell(r, 6).value)})

# Component Categories (draft) for the required-systems warning.
cats = []
ws = wb["Component Categories"]
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 1).value: cats.append({"component": text(ws.cell(r, 1).value), "required": text(ws.cell(r, 2).value) == "Yes",
        "spaces": num(ws.cell(r, 3).value), "description": text(ws.cell(r, 4).value), "status": text(ws.cell(r, 5).value)})

ids = [i["_id"] for i in items]
assert len(ids) == len(set(ids)), "duplicate ids"
json.dump({"exportedAt": raw.get("exportedAt"), "overrides": OVERRIDES, "items": items, "mishapRef": mishap_ref, "categories": cats},
          open(OUT, "w"), ensure_ascii=False)
from collections import Counter
print("overrides", len(OVERRIDES))
print("items", len(items), dict(Counter(i["flags"][MOD]["component"]["slot"] for i in items)))
print("stations", [(i["name"], i["flags"][MOD]["component"]["stationRoles"]) for i in items if i["flags"][MOD]["component"].get("stationRoles")][:40])
